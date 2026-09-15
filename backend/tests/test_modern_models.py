from __future__ import annotations

import asyncio
import base64
import io
import json
import wave
from pathlib import Path
from types import SimpleNamespace

import httpx
import pytest
from openai import OpenAI
from PIL import Image

from app import config
from app.llm import modern_media, multimodal_gateway, openai_client, spark_client, ark_video
from app.services.native_ppt import Outline, render


@pytest.fixture
def transport(monkeypatch):
    monkeypatch.setenv("LLM_PROVIDER", "openai")
    monkeypatch.setenv("DEMO_MODE", "false")
    monkeypatch.setenv("OPENAI_API_KEY", "test-only-key")
    monkeypatch.setenv("OPENAI_BASE_URL", "https://api.example.test/v1")
    monkeypatch.setenv("OPENAI_TEXT_MODEL", "gpt-6-astra")
    monkeypatch.setenv("OPENAI_IMAGE_MODEL", "gpt-image-2.5-sunburst")
    config.get_settings.cache_clear()
    calls = []
    def install(handler):
        def record(request):
            calls.append(request)
            return handler(request)
        monkeypatch.setattr(openai_client, "client", lambda: OpenAI(api_key="test-only-key",
            base_url="https://api.example.test/v1", max_retries=0,
            http_client=httpx.Client(transport=httpx.MockTransport(record))))
        return calls
    return install


def response(text="先左，再根，最后右。", status="completed"):
    return {"id":"resp_test", "object":"response", "status":status, "created_at":1,
            "model":"gpt-6-astra", "output":[{"id":"msg_test", "type":"message",
            "role":"assistant", "status":"completed", "content":[
                {"type":"output_text", "text":text, "annotations":[]}]}]}


def test_responses_sdk_contract_and_json_arrays(transport):
    calls = transport(lambda _: httpx.Response(200, json=response('[{"kp":"tree"}]')))
    result = asyncio.run(spark_client.llm_complete("JSON array", system="course rules", json_mode=True, max_tokens=700))
    assert json.loads(result) == [{"kp":"tree"}]
    request = json.loads(calls[0].content)
    assert calls[0].url.path == "/v1/responses"
    assert request["model"] == "gpt-6-astra" and request["store"] is False
    assert request["input"][0]["role"] == "system"
    assert request["reasoning"]["effort"] == "low"
    assert "temperature" not in request


@pytest.mark.parametrize("status,text", [("incomplete","partial"),("completed","")])
def test_incomplete_answers_never_become_mock_lessons(transport, status, text):
    transport(lambda _: httpx.Response(200,json=response(text,status)))
    with pytest.raises(RuntimeError, match="OpenAI 调用失败"):
        asyncio.run(spark_client.llm_complete("course", allow_mock_fallback=True))


def test_provider_errors_do_not_expose_body_or_retry(transport):
    calls=transport(lambda _: httpx.Response(429,json={"error":{"message":"secret-key-and-private-prompt",
                    "type":"insufficient_quota","code":"credit_balance_exhausted"}}))
    with pytest.raises(RuntimeError) as caught:
        asyncio.run(spark_client.llm_complete("private prompt"))
    assert "secret-key" not in str(caught.value) and "private" not in str(caught.value)
    assert "429" in str(caught.value) and len(calls)==1


def test_stream_partial_failure_does_not_append_a_different_model(transport):
    events = [
        {"type":"response.output_text.delta", "delta":"左子树", "item_id":"x", "output_index":0,"content_index":0,"sequence_number":1},
        {"type":"response.incomplete", "response":response("", "incomplete"), "sequence_number":2},
    ]
    data="".join("data: "+json.dumps(e)+"\n\n" for e in events)
    calls=transport(lambda _:httpx.Response(200,content=data,headers={"content-type":"text/event-stream"}))
    collected=[]
    async def read():
        async for piece in spark_client.llm_stream("inorder"):
            collected.append(piece)
    with pytest.raises(RuntimeError): asyncio.run(read())
    assert collected==["左子树"] and len(calls)==1


def test_image_sdk_and_cache_separate_tenants_models_and_endpoint(transport, monkeypatch, tmp_path):
    buffer=io.BytesIO();Image.new("RGB",(128,128),"green").save(buffer,format="PNG")
    calls=transport(lambda _:httpx.Response(200,json={"created":1,"data":[{"b64_json":base64.b64encode(buffer.getvalue()).decode()}]}))
    monkeypatch.setattr(modern_media,"GEN_DIR",tmp_path)
    first=modern_media.generate_image("inorder",1536,864,"learner-a")
    assert modern_media.generate_image("inorder",1536,864,"learner-a")["cached"] is True
    second=modern_media.generate_image("inorder",1536,864,"learner-b")
    assert first["image_url"] != second["image_url"] and len(calls)==2
    assert json.loads(calls[0].content)["model"]=="gpt-image-2.5-sunburst"
    monkeypatch.setenv("OPENAI_IMAGE_MODEL","gpt-image-2.5-flare");config.get_settings.cache_clear()
    assert modern_media.generate_image("inorder",1536,864,"learner-a")["image_url"] != first["image_url"]
    monkeypatch.setenv("OPENAI_BASE_URL","https://another.example.test/v1");config.get_settings.cache_clear()
    assert modern_media.generate_image("inorder",1536,864,"learner-a")["cached"] is False


def test_invalid_image_response_is_not_persisted(transport, monkeypatch, tmp_path):
    transport(lambda _:httpx.Response(200,json={"created":1,"data":[{"b64_json":base64.b64encode(b'not an image').decode()}]}))
    monkeypatch.setattr(modern_media,"GEN_DIR",tmp_path)
    with pytest.raises(Exception): modern_media.generate_image("inorder",512,512,"learner-a")
    assert list(tmp_path.iterdir())==[]


def test_pcm_is_wrapped_in_wav_before_transcription(monkeypatch):
    captured={}
    def create(**kwargs): captured.update(kwargs);return SimpleNamespace(text="中序遍历")
    class Client:
        audio=SimpleNamespace(transcriptions=SimpleNamespace(create=create))
        def __enter__(self):return self
        def __exit__(self,*args):pass
    monkeypatch.setattr(openai_client,"client",Client)
    data=modern_media.transcribe(b'\x00\x01'*1600,"raw",16000)
    with wave.open(io.BytesIO(captured["file"][1]),"rb") as wav:
        assert wav.getframerate()==16000 and wav.getnframes()==1600 and wav.getsampwidth()==2
    assert data["text"]=="中序遍历" and data["model"]=="gpt-transcribe"


def test_demo_mode_never_contacts_media_provider(monkeypatch):
    monkeypatch.setenv("IMAGE_PROVIDER","openai");config.get_settings.cache_clear()
    def fail(*args):raise AssertionError("External call in demo mode")
    monkeypatch.setattr(modern_media,"generate_image",fail)
    result=asyncio.run(multimodal_gateway.text_to_image("tree"))
    assert result["ok"] is False


def test_native_deck_is_editable_and_schema_rejects_extra_data(tmp_path):
    from pptx import Presentation
    spec={"title":"二叉树复习", "slides":[{"title":"中序遍历", "bullets":["左子树、根结点、右子树"]}]*3}
    path=tmp_path/"lesson.pptx";render(Outline.model_validate(spec),path)
    deck=Presentation(path)
    assert len(deck.slides)==3
    assert any("左子树" in shape.text for shape in deck.slides[0].shapes if shape.has_text_frame)
    with pytest.raises(ValueError): Outline.model_validate({**spec,"speaker_script":"unvalidated"})


def test_native_deck_resolves_citations_and_rejects_fabricated_ones(monkeypatch, tmp_path):
    from pptx import Presentation
    from app.services import native_ppt
    spec={"title":"二叉树复习", "slides":[{"title":"中序遍历", "bullets":["左子树、根结点、右子树。[^1]"]} for _ in range(3)]}
    async def complete(*args, **kwargs):
        return json.dumps(spec, ensure_ascii=False)
    monkeypatch.setattr(native_ppt,"llm_complete",complete)
    monkeypatch.setattr(native_ppt,"is_demo",lambda:False)
    monkeypatch.setattr(native_ppt,"GEN_DIR",tmp_path)
    result=asyncio.run(native_ppt.create("按给定资料复习", citations=[{"index":1,"citation":"课程讲义 · 第6章 · 第1段"}]))
    deck=Presentation(tmp_path/Path(result["url"]).name)
    assert result["content_pages"]==3 and result["reference_pages"]==1 and result["total_pages"]==4
    assert any("[^1] 课程讲义 · 第6章 · 第1段" in shape.text for shape in deck.slides[-1].shapes if shape.has_text_frame)
    spec["slides"][0]["bullets"]=["无法核验的内容。[^2]"]
    with pytest.raises(ValueError, match="without a supplied source"):
        asyncio.run(native_ppt.create("按给定资料复习", citations=[{"index":1,"citation":"课程讲义"}]))
    assert len(list(tmp_path.glob("*.pptx")))==1


def test_ark_contract_validation_and_terminal_states(monkeypatch):
    monkeypatch.setenv("ARK_API_KEY","test-only");config.get_settings.cache_clear()
    calls=[]
    def request(method,url,**kwargs):
        calls.append((method,url,kwargs))
        return SimpleNamespace(ok=True,json=lambda: {"id":"cgt-test"} if method=="POST" else
              {"status":"succeeded","model":"doubao-seedance-2-5-260628","content":{"video_url":"https://assets.example.test/a.mp4"}})
    monkeypatch.setattr(ark_video.requests,"request",request)
    task=ark_video.create_task("A binary tree")
    assert task=="ark_cgt-test"
    assert calls[0][2]["json"]["generate_audio"] is False
    assert calls[0][2]["json"]["watermark"] is True
    assert calls[0][1].endswith("/api/v3/contents/generations/tasks")
    assert ark_video.poll_task(task)["provider"]=="ark"
    with pytest.raises(ValueError):ark_video.poll_task("ark_../../private")
    assert len(calls)==2
