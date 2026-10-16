from __future__ import annotations

from types import SimpleNamespace

import numpy as np

from app.agents import doc_agent, media_agent, mindmap_agent, quiz_agent
from app.rag import ingest, retriever


def test_word_document_text_and_table_are_extractable(tmp_path):
    from docx import Document

    path = tmp_path / "lesson.docx"
    document = Document()
    document.add_heading("排序基础", level=1)
    document.add_paragraph("插入排序把当前元素放进前面已有序的区域。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "稳定性"
    table.cell(0, 1).text = "相等元素相对次序不变"
    document.save(path)

    text = ingest._read_docx(path)  # noqa: SLF001

    assert "插入排序" in text
    assert "稳定性 | 相等元素相对次序不变" in text


def test_course_and_personal_sources_can_be_retrieved_together(monkeypatch):
    course_hit = {"text": "课程片段", "source": retriever.COURSE_SOURCE, "source_id": "sorting"}
    personal_hit = {"text": "个人片段", "source": "个人笔记", "source_id": "source-a"}
    store = SimpleNamespace(search=lambda *args, **kwargs: [course_hit])
    monkeypatch.setattr(retriever, "get_store", lambda: store)
    monkeypatch.setattr(retriever, "embed_texts", lambda items, **kwargs: np.ones((len(items), 2)))
    monkeypatch.setattr(retriever, "_hits_from_sources", lambda *args, **kwargs: [personal_hit])
    monkeypatch.setattr(retriever, "rerank", lambda query, hits, top_k: (hits[:top_k], "test"))

    hits = retriever.retrieve(
        "排序",
        source_ids=[ingest.COURSE_SCOPE_ID, "source-a"],
        user_id="alice",
    )

    assert {hit["source_id"] for hit in hits} == {"sorting", "source-a"}


def test_pdf_exam_dump_is_cleaned_before_display_and_narration():
    raw = "（7）若一组记录的排序码为…… A. 38,40 B. 40,38 答案：B 解析：快速排序。"
    cleaned = doc_agent._clean_source_excerpt(raw)  # noqa: SLF001

    assert "答案" not in cleaned
    assert "解析" not in cleaned
    assert "A." not in cleaned


def test_quiz_quality_gate_rejects_stems_that_leak_answers():
    question = quiz_agent._validate({  # noqa: SLF001
        "type": "design",
        "stem": "请说明快速排序。答案：先进行划分。",
        "answer": "先进行划分",
    }, "quick_sort")

    assert question is None


def test_framework_map_fills_sparse_source_points_to_five_branches():
    points = mindmap_agent._complete_learning_branches(  # noqa: SLF001
        ["排序用于重新排列元素"],
        has_sources=True,
    )

    assert len(points) == 5
    assert any("关键步骤" in point for point in points)
    assert any("易错边界" in point for point in points)


def test_video_script_preserves_short_lesson_without_pasting_source_excerpts():
    narration = "排序就是按规则重新排列。把 3、1、2 从小到大排列，得到 1、2、3。元素保持不变。"
    scenes = [
        {"t": "0-10s", "visual": "摆出 3、1、2", "caption": "输入"},
        {"t": "10-25s", "visual": "移成 1、2、3", "caption": "升序"},
    ]
    script = media_agent._normalize_script(  # noqa: SLF001
        {"title": "排序", "narration": narration, "scenes": scenes},
        [{"text": "插入排序逐步扩大有序区域，每一步都把当前元素插入正确位置。"} for _ in range(5)],
        "排序",
    )

    assert script["scenes"] == scenes
    assert script["target_seconds"] == 25
    assert script["video_clip_seconds"] == 10
    assert script["narration"] == narration
    assert "插入排序逐步扩大有序区域" not in script["narration"]


def test_video_narration_removes_markdown_source_syntax():
    cleaned = media_agent._clean_script_text(  # noqa: SLF001
        "## 基础排序\n插入排序把当前元素放到正确位置；不要展示 `source code`。"
    )

    assert "##" not in cleaned
    assert "`" not in cleaned
    assert "基础排序" in cleaned


def test_video_normalization_cannot_disguise_wrong_topic_by_appending_its_name():
    prefix = "先看模型原始讲解的第一句话。"
    narration = prefix + "这段正文没有写出课程主题。" * 50
    script = media_agent._normalize_script(  # noqa: SLF001
        {"title": "演示", "narration": narration, "scenes": []},
        [],
        "图的遍历",
    )

    assert script["narration"] == narration
    assert media_agent._script_on_topic(script, "图的遍历", "graph_traverse", []) is False  # noqa: SLF001


def test_video_storyboard_repairs_missing_timing_without_discarding_steps():
    script = media_agent._normalize_script(  # noqa: SLF001
        {
            "narration": "二叉树用于表达层次关系。",
            "target_seconds": 90,
            "scenes": [{"visual": str(index), "caption": str(index)} for index in range(9)],
        },
        [],
        "二叉树",
    )

    assert [scene["visual"] for scene in script["scenes"]] == [str(index) for index in range(9)]
    assert [scene["t"] for scene in script["scenes"]] == [
        f"{index * 10}-{(index + 1) * 10}s" for index in range(9)
    ]
    assert script["target_seconds"] == 90


def test_video_topic_gate_rejects_adjacent_course_topic():
    script = {
        "narration": "图的遍历从一个顶点出发；接着详细讲最短路径的松弛操作。",
        "scenes": [{"visual": "图的遍历", "caption": "最短路径"}],
    }

    assert media_agent._script_on_topic(script, "图的遍历", "graph_traverse", []) is False  # noqa: SLF001


def test_video_topic_gate_accepts_mixed_language_label_with_example():
    script = {
        "narration": "Python 里的变量，就是给一个值贴上名字。把名字 x 贴到 3 上叫赋值。",
        "scenes": [{"visual": "名字 x 贴到数值 3 上", "caption": "一个等号做赋值"}],
    }

    for label in ("Python 变量与赋值：x = 3", "Python变量与赋值: x=3"):
        assert media_agent._script_on_topic(script, label, "custom_topic", []) is True  # noqa: SLF001


def test_video_topic_gate_cannot_pass_on_language_name_alone():
    script = {
        "narration": "Python 的循环可以重复做一件事。",
        "scenes": [{"visual": "重复运行三个步骤", "caption": "循环"}],
    }

    assert media_agent._script_on_topic(script, "Python 变量与赋值：x = 3", "custom_topic", []) is False  # noqa: SLF001
