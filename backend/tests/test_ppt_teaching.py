from __future__ import annotations

import pytest
from pptx import Presentation
from app.agents.doc_agent import _ppt_query
from app.services.native_ppt import Outline, Slide, render, validate_teaching, _wrap


def lesson():
    notes = "函数值取决于函数在该点的定义，极限考察的是自变量靠近该点时的变化趋势。判断连续必须同时检查函数值、极限及二者的关系，不能只算出极限就结束。"
    return Outline.model_validate({"title": "极限与连续", "slides": [
        {"title": "连续的条件", "bullets": ["函数在该点有定义，极限存在，且极限等于函数值。"], "notes": notes},
        {"title": "函数值与极限", "layout": "comparison", "table": {"headers": ["对象", "判断方法"], "rows": [["函数值", "查看该点的定义"], ["极限", "比较左右两侧趋近的结果"]]}, "notes": notes},
        {"title": "可去间断点", "layout": "example", "question": "f(x)=(x²−1)/(x−1)，判断 x=1 处的连续性。", "bullets": ["x=1 时分母为零，函数没有定义。", "x≠1 时可约分为 x+1，趋近1时结果趋近2。"], "answer": "该点为可去间断点，补定义 f(1)=2 后连续。", "notes": notes},
        {"title": "独立练习", "layout": "practice", "question": "f(x) 在 x≠0 时等于1，在 x=0 时等于2，是否连续？", "answer": "不连续。左右极限都等于1，但 f(0)=2，不满足极限等于函数值。", "notes": notes},
    ]})


def test_new_generation_rejects_a_bullet_only_outline():
    old = Outline.model_validate({"title": "复习", "slides": [{"title": "概念", "bullets": ["定义"]}] * 3})
    with pytest.raises(ValueError, match="完整例题.*独立自测.*讲解备注"):
        validate_teaching(old)
    validate_teaching(lesson())


def test_lesson_preserves_editable_table_notes_and_separate_answers(tmp_path):
    target = tmp_path / "lesson.pptx"
    render(lesson(), target)
    deck = Presentation(target)
    assert len(deck.slides) == 5
    tables = [shape.table for slide in deck.slides for shape in slide.shapes if shape.has_table]
    assert len(tables) == 1
    assert tables[0].cell(2, 0).text == "极限"
    assert all("不能只算出极限" in slide.notes_slide.notes_text_frame.text for slide in deck.slides)
    pages = ["\n".join(shape.text for shape in slide.shapes if shape.has_text_frame) for slide in deck.slides]
    question_page = next(i for i, text in enumerate(pages) if "独立练习" in text)
    assert "不连续" not in pages[question_page]
    assert "不连续" in pages[question_page + 1]
    assert "第 1 步" in pages[2] and "补定义" in pages[2]


def test_long_content_continues_without_losing_conditions_or_covering_footer(tmp_path):
    spec = lesson()
    long_point = "判断时先检查左右极限是否都存在，随后比较两者是否相等。" * 5 + "必须保留末尾条件。"
    spec.slides[0].bullets = [long_point] * 4
    target = tmp_path / "long.pptx"
    render(spec, target, ["[^1] " + "完整来源与章节说明" * 30])
    deck = Presentation(target)
    assert len(deck.slides) > 5
    texts = []
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text.replace("\v", ""))
                assert shape.left >= 0 and shape.top >= 0
                assert shape.left + shape.width <= deck.slide_width
                assert shape.top + shape.height <= deck.slide_height
                if shape.top / 914400 > 2 and shape.top / 914400 < 6.9:
                    assert (shape.top + shape.height) / 914400 < 6.9
    assert "\n".join(texts).count("必须保留末尾条件。") == 4
    assert "完整来源与章节说明" * 30 in "\n".join(texts)


def test_example_and_comparison_require_complete_fields():
    with pytest.raises(ValueError, match="题目和答案"):
        Slide(title="例题", layout="example", bullets=["算出极限", "写结论"])
    with pytest.raises(ValueError, match="match the headers"):
        Slide(title="对比", layout="comparison", table={"headers": ["对象", "条件"], "rows": [["一"], ["二", "完整"]]})


def test_learning_conversation_does_not_cut_off_source_conditions():
    query = _ppt_query("极限", "复习", "零基础", "完整资料条件：左右极限都存在且相等。", source_bound=True,
                       conversation="讨论这个例子的理由。" * 1300)
    assert "左右极限都存在且相等。" in query
    assert "只允许使用下方所选资料" in query
    assert "不作为课程来源" in query


def test_long_comparison_repeats_headers_and_preserves_every_cell(tmp_path):
    spec = lesson()
    table = spec.slides[1].table
    table.rows = [[f"对象{i}", "只有左右极限都存在并且相等时才能进一步比较函数值。" * 3] for i in range(6)]
    target = tmp_path / "table.pptx"
    render(spec, target)
    tables = [shape.table for slide in Presentation(target).slides for shape in slide.shapes if shape.has_table]
    assert len(tables) > 1
    rows = []
    for rendered in tables:
        assert [cell.text for cell in rendered.rows[0].cells] == table.headers
        rows.extend([[cell.text.replace("\v", "") for cell in row.cells] for row in list(rendered.rows)[1:]])
        assert sum(row.height for row in rendered.rows) / 914400 <= 4.4
        assert rendered.cell(1, 0).margin_left > 0
        assert rendered.cell(1, 0).text_frame.margin_left > 0
    assert rows == table.rows


def test_answer_continuation_keeps_a_long_title_within_two_lines(tmp_path):
    spec = lesson()
    spec.slides[-1].title = "复习函数极限连续性及其判断条件与易错概念之间的区别" + "题" * 14
    spec.slides[-1].question = "判断连续之前必须先检查函数在该点是否有定义。" * 12
    spec.slides[-1].answer = "先判断左右极限是否都存在且相等，再比较函数值。" * 12
    target = tmp_path / "title.pptx"
    render(spec, target)
    titles = [slide.shapes[1] for slide in Presentation(target).slides]
    assert any("自测解析" in shape.text and "续" in shape.text for shape in titles)
    assert all(len(shape.text.split("\v")) <= 2 for shape in titles)


def test_wrapping_keeps_short_math_tokens_and_closing_punctuation_together():
    text = "先检查函数值" * 4 + " f(x)=1，然后比较左右极限 lim(x→a⁻) 与 lim(x→a⁺)。"
    lines = _wrap(text, 5, 22)
    assert "".join(lines) == text
    assert any("f(x)=1" in line for line in lines)
    assert any("lim(x→a⁻)" in line for line in lines)
    assert any("lim(x→a⁺)" in line for line in lines)
    assert not any(line.startswith(("，", "。")) for line in lines)
