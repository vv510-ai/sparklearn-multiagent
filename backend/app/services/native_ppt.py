"""Editable study decks with complete explanations, examples and self-checks."""
from __future__ import annotations

import asyncio
import re
import unicodedata
import uuid
from pathlib import Path
from typing import Annotated, Literal

from pydantic import BaseModel, ConfigDict, Field, model_validator

from ..config import GEN_DIR, get_settings, is_demo
from .model_preferences import resolve
from ..llm.spark_client import active_llm_label, active_provider, llm_complete, parse_json
from ..prompt_templates import invoke_learning_prompt


class Comparison(BaseModel):
    model_config = ConfigDict(extra="forbid")
    headers: list[Annotated[str, Field(min_length=1, max_length=24)]] = Field(min_length=2, max_length=3)
    rows: list[list[Annotated[str, Field(min_length=1, max_length=100)]]] = Field(min_length=2, max_length=6)

    @model_validator(mode="after")
    def rectangular(self):
        if any(len(row) != len(self.headers) for row in self.rows):
            raise ValueError("PPT comparison rows must match the headers")
        return self


class Slide(BaseModel):
    model_config = ConfigDict(extra="forbid")
    title: str = Field(min_length=1, max_length=40)
    layout: Literal["concept", "example", "comparison", "practice", "summary"] = "concept"
    bullets: list[Annotated[str, Field(min_length=1, max_length=160)]] = Field(default_factory=list, max_length=4)
    question: str = Field(default="", max_length=400)
    answer: str = Field(default="", max_length=500)
    notes: str = Field(default="", max_length=1600)
    table: Comparison | None = None

    @model_validator(mode="after")
    def complete_content(self):
        if self.layout in {"example", "practice"} and not (self.question.strip() and self.answer.strip()):
            raise ValueError("例题与自测必须同时提供题目和答案")
        if self.layout == "example" and len(self.bullets) < 2:
            raise ValueError("例题必须包含至少两步完整解法")
        if self.layout == "comparison" and self.table is None:
            raise ValueError("对照页必须提供表格")
        if not (self.bullets or self.question.strip() or self.table):
            raise ValueError("PPT 页面不能只有标题")
        return self


class Outline(BaseModel):
    model_config = ConfigDict(extra="forbid")
    title: str = Field(min_length=1, max_length=60)
    slides: list[Slide] = Field(min_length=3, max_length=10)


def validate_teaching(outline: Outline) -> None:
    """A list of headings is valid legacy data, but not a finished new lesson."""
    missing = []
    if not any(slide.layout == "example" for slide in outline.slides):
        missing.append("有题目、步骤和结论的完整例题")
    if not any(slide.layout == "practice" for slide in outline.slides):
        missing.append("带答案的独立自测")
    if any(len(slide.notes.strip()) < 40 for slide in outline.slides):
        missing.append("逐页讲解备注")
    if missing:
        raise ValueError("PPT 内容不完整，缺少：" + "、".join(missing))


def _wrap(text: str, width: float, size: float) -> list[str]:
    """Conservative CJK-aware wrapping; never discard overflowing content."""
    capacity = max(8, int(width * 72 / size * 1.75))
    lines = []
    for paragraph in text.splitlines() or [""]:
        line, used = "", 0
        for char in paragraph:
            units = 2 if unicodedata.east_asian_width(char) in {"W", "F", "A"} else 1
            if line and used + units > capacity:
                lines.append(line)
                line, used = "", 0
            line += char
            used += units
        lines.append(line)
    return lines


def _text_pages(title: str, blocks: list[str], notes: str, *, size: int = 22) -> list[dict]:
    pages, current, used = [], [], 0
    for block in blocks:
        lines = _wrap(block, 11.85, size)
        # At most 11 body lines plus paragraph spacing fit above the footer.
        for offset in range(0, len(lines), 10):
            part = lines[offset:offset + 10]
            cost = len(part) + 0.55
            if current and used + cost > 11.5:
                pages.append({"title": title, "blocks": current, "notes": notes, "size": size})
                current, used = [], 0
            current.append(part)
            used += cost
    if current:
        pages.append({"title": title, "blocks": current, "notes": notes, "size": size})
    for index, page in enumerate(pages):
        if index:
            page["title"] += f"（续 {index}）"
    return pages


def _table_pages(slide: Slide) -> list[dict]:
    table = slide.table
    assert table is not None
    width, size = 11.85 / len(table.headers), 19
    def height(row):
        return max(len(_wrap(cell, width - .28, size)) for cell in row) * size * 1.3 / 72 + .22
    header_height = height(table.headers)
    pages, rows, heights = [], [], []
    def append_page():
        pages.append({"title": slide.title, "headers": table.headers, "rows": list(rows),
                      "heights": [header_height, *heights], "notes": slide.notes, "size": size})
    for row in table.rows:
        row_height = height(row)
        if header_height + row_height > 4.4:
            raise ValueError("PPT 对照单元格过长，请拆成更短的比较项")
        if rows and header_height + sum(heights) + row_height > 4.4:
            append_page()
            rows, heights = [], []
        rows.append(row)
        heights.append(row_height)
    if rows:
        append_page()
    for index, page in enumerate(pages):
        if index:
            page["title"] += f"（续 {index}）"
    return pages


def _pages(outline: Outline, references: list[str]) -> list[dict]:
    pages = []
    for item in outline.slides:
        blocks = []
        if item.question:
            blocks.append("题目：" + item.question)
        for index, bullet in enumerate(item.bullets, 1):
            prefix = f"第 {index} 步：" if item.layout == "example" else ""
            blocks.append(prefix + bullet)
        if item.answer and item.layout != "practice":
            blocks.append("结论：" + item.answer)
        pages.extend(_text_pages(item.title, blocks, item.notes))
        if item.table:
            pages.extend(_table_pages(item))
        if item.layout == "practice":
            pages.extend(_text_pages("自测解析：" + item.title,
                                     ["题目：" + item.question, "答案：" + item.answer], item.notes))
    if references:
        pages.extend(_text_pages("资料来源与引用编号", references, "\n".join(references), size=19))
    return pages


def render(outline: Outline, target: Path, references: list[str] | None = None) -> int:
    # Keep the deployed renderer and editable Office objects; no image service is needed.
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    from pptx.oxml.xmlchemy import OxmlElement
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    pages = _pages(outline, references or [])
    for index, item in enumerate(pages):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("F4F7F4")
        def style(frame, text, size, color="263E37", bold=False):
            frame.clear()
            frame.word_wrap = False
            frame.margin_left = frame.margin_right = 0
            frame.margin_top = frame.margin_bottom = 0
            paragraph = frame.paragraphs[0]
            paragraph.text = text
            paragraph.font.name = "Microsoft YaHei"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
            paragraph.line_spacing = Pt(size * 1.3)
            properties = paragraph._p.get_or_add_pPr().get_or_add_defRPr()
            east_asian = OxmlElement("a:ea")
            east_asian.set("typeface", "Microsoft YaHei")
            properties.append(east_asian)
        def box(text, x, y, width, height, size, color="263E37", bold=False):
            frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height)).text_frame
            style(frame, text, size, color, bold)
        box("知行循证 / 课程复习", .7, .35, 11.85, .4, 13, "167259")
        title = " ".join(item["title"].splitlines())
        title_size = 32 if len(_wrap(title, 11.85, 32)) <= 2 else 28
        box("\v".join(_wrap(title, 11.85, title_size)), .7, .88, 11.85, 1.18, title_size, "14352E", True)
        if "rows" in item:
            height = sum(item["heights"])
            table = slide.shapes.add_table(len(item["rows"]) + 1, len(item["headers"]),
                                          Inches(.7), Inches(2.16), Inches(11.85), Inches(height)).table
            for row_index, row in enumerate([item["headers"], *item["rows"]]):
                table.rows[row_index].height = Inches(item["heights"][row_index])
                for col_index, text in enumerate(row):
                    cell = table.cell(row_index, col_index)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string("167259" if row_index == 0 else "FFFFFF")
                    wrapped = "\v".join(_wrap(text, 11.85 / len(row) - .28, item["size"]))
                    style(cell.text_frame, wrapped, item["size"], "FFFFFF" if row_index == 0 else "263E37", row_index == 0)
                    cell.margin_left = cell.margin_right = Inches(.14)
                    cell.margin_top = cell.margin_bottom = Inches(.11)
        else:
            y, size = 2.18, item["size"]
            for lines in item["blocks"]:
                height = len(lines) * size * 1.3 / 72 + .07
                box("\v".join(lines), .7, y, 11.85, height, size)
                y += height + .15
        slide.notes_slide.notes_text_frame.text = item["notes"]
        box("可编辑课件 · 详细讲解见演讲者备注", .7, 7.02, 10.4, .3, 11, "63766E")
        box(f"{index + 1:02d} / {len(pages):02d}", 11.7, 7.02, 1, .3, 11, "63766E")
    temp = target.with_suffix(".part")
    try:
        deck.save(temp)
        temp.replace(target)
    finally:
        temp.unlink(missing_ok=True)
    return len(pages)


async def create(query: str, *, citations: list[dict] | None = None) -> dict:
    if is_demo() and resolve("CAPABILITY=media.ppt_outline") is None:
        raise RuntimeError("Editable AI slides require a configured text model")
    call = invoke_learning_prompt("media.ppt_outline", content=query)
    raw = await llm_complete(call.prompt, system=call.system, role="reasoner",
                             json_mode=True, json_object=True, max_tokens=8500,
                             reasoning_mode="disabled", temperature=0.2,
                             allow_model_fallback=False, allow_mock_fallback=False)
    outline = Outline.model_validate(parse_json(raw))
    sources = {}
    for position, citation in enumerate(citations or [], start=1):
        index = int(citation.get("index", position))
        label = str(citation.get("citation") or " · ".join(str(citation[k]) for k in ("source", "chapter", "page") if citation.get(k))).strip()
        if index > 0 and label:
            sources[index] = label
    used = {int(index) for index in re.findall(r"\[\^(\d+)\]", outline.model_dump_json())}
    if used - sources.keys():
        raise ValueError("PPT contains a citation without a supplied source")
    validate_teaching(outline)
    references = [f"[^{index}] {label}" for index, label in sorted(sources.items())]
    target = GEN_DIR / ("lesson_" + uuid.uuid4().hex + ".pptx")
    target.parent.mkdir(parents=True, exist_ok=True)
    total_pages = await asyncio.to_thread(render, outline, target, references)
    content_pages = len(_pages(outline, []))
    return {"status": "ready", "title": outline.title, "total_pages": total_pages,
            "content_pages": content_pages, "reference_pages": total_pages - content_pages,
            "url": "/static/gen/" + target.name, "provider": "native_pptx",
            "model": get_settings().deepseek_model if active_provider() == "deepseek" else active_llm_label()}
