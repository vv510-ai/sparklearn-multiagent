"""Build the lightweight RAG knowledge base.

Default course notes live in data/seed_corpus/*.md. Extra sources can be added
under data/extra_sources as .md, .txt, or .json files. All sources share one
vector store and carry citation metadata.
"""
from __future__ import annotations

import hashlib
import json
import logging
import re
from pathlib import Path
from typing import Any

from ..config import CORPUS_DIR, DATA_DIR, EXTRA_SOURCES_DIR, UPLOAD_SOURCES_DIR
from .embedding import embed_texts
from .vector_store import get_store

log = logging.getLogger("sparklearn.ingest")

CHUNK = 420
OVERLAP = 50
MANIFEST = DATA_DIR / "vector_store_manifest.json"
SUPPORTED_EXTRA = {".md", ".txt", ".json"}
SUPPORTED_UPLOAD = {".md", ".txt", ".pdf", ".docx", ".png", ".jpg", ".jpeg"}
IMAGE_UPLOAD_SUFFIXES = {".png", ".jpg", ".jpeg"}
COURSE_SCOPE_ID = "__course__"
SOURCE_STATUSES = {"processing", "indexed", "degraded", "failed", "quarantined", "inactive"}
SELECTABLE_SOURCE_STATUSES = {"indexed", "degraded"}
INACTIVE_SOURCE_STATUSES = {"failed", "quarantined", "inactive"}


def _front_matter(text: str) -> tuple[dict[str, str], str]:
    meta = {
        "source": "未知讲义",
        "chapter": "",
        "kp": "",
        "license": "",
        "url": "",
        "type": "markdown",
        "tags": "",
    }
    m = re.match(r"^---\n(.*?)\n---\n", text, flags=re.S)
    if m:
        for line in m.group(1).splitlines():
            if ":" in line:
                k, v = line.split(":", 1)
                meta[k.strip()] = v.strip()
        text = text[m.end():]
    return meta, text


def _split(text: str) -> list[str]:
    protected: list[str] = []

    def _hold(m: re.Match) -> str:
        protected.append(m.group(0))
        return f"\x00CODE{len(protected) - 1}\x00"

    text = re.sub(r"```.*?```", _hold, text, flags=re.S)
    sections = re.split(r"(?=^##\s)", text, flags=re.M)
    chunks: list[str] = []
    for sec in sections:
        sec = sec.strip()
        if not sec:
            continue
        if len(sec) <= CHUNK:
            chunks.append(sec)
        else:
            i = 0
            while i < len(sec):
                chunks.append(sec[i:i + CHUNK])
                i += CHUNK - OVERLAP

    out: list[str] = []
    for c in chunks:
        def _restore(m: re.Match) -> str:
            return protected[int(m.group(1))]

        out.append(re.sub(r"\x00CODE(\d+)\x00", _restore, c).strip())
    return [c for c in out if len(c) > 30]


def _read_text(path: Path) -> str:
    for enc in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def _source_file_signature(path: Path) -> dict[str, Any]:
    data = path.read_bytes()
    try:
        rel = path.relative_to(DATA_DIR).as_posix()
    except ValueError:
        rel = str(path)
    return {"path": rel, "sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data)}


def _current_manifest() -> dict[str, Any]:
    files = sorted(CORPUS_DIR.glob("*.md"))
    if EXTRA_SOURCES_DIR.exists():
        files += sorted(
            p for p in EXTRA_SOURCES_DIR.rglob("*")
            if p.is_file() and p.suffix.lower() in SUPPORTED_EXTRA
        )
    if UPLOAD_SOURCES_DIR.exists():
        files += sorted(
            p for p in UPLOAD_SOURCES_DIR.rglob("*")
            if p.is_file()
            and (p.suffix.lower() in SUPPORTED_UPLOAD or p.name.endswith(".meta.json"))
        )
    return {"version": 3, "files": [_source_file_signature(p) for p in files]}


def _manifest_matches(current: dict[str, Any]) -> bool:
    if not MANIFEST.exists():
        return False
    try:
        return json.loads(MANIFEST.read_text(encoding="utf-8")) == current
    except Exception:
        return False


def _write_manifest(manifest: dict[str, Any]) -> None:
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")


def _tags(value: Any) -> list[str]:
    if isinstance(value, list):
        return [str(x).strip() for x in value if str(x).strip()]
    if isinstance(value, str):
        return [x.strip() for x in value.split(",") if x.strip()]
    return []


def _chunk_record(
    *,
    text: str,
    source: str,
    chapter: str,
    kp: str,
    page: int,
    source_type: str,
    url: str = "",
    tags: list[str] | None = None,
    source_id: str = "",
    user_id: str = "",
) -> dict[str, Any]:
    return {
        "text": text,
        "source": source,
        "chapter": chapter,
        "kp": kp,
        "page": page,
        "source_type": source_type,
        "url": url,
        "tags": tags or [],
        "source_id": source_id,
        "user_id": user_id,
    }


def _load_seed_corpus() -> list[dict[str, Any]]:
    all_chunks: list[dict[str, Any]] = []
    for path in sorted(CORPUS_DIR.glob("*.md")):
        meta, body = _front_matter(_read_text(path))
        for page, piece in enumerate(_split(body), start=1):
            all_chunks.append(_chunk_record(
                text=piece,
                source=meta.get("source") or path.stem,
                chapter=meta.get("chapter") or path.stem,
                kp=meta.get("kp", ""),
                page=page,
                source_type=meta.get("type") or "lecture",
                url=meta.get("url", ""),
                tags=_tags(meta.get("tags")),
                source_id=path.stem,
            ))
    return all_chunks


def _json_items(path: Path) -> list[dict[str, Any]]:
    raw = json.loads(_read_text(path))
    if isinstance(raw, list):
        return [x for x in raw if isinstance(x, dict)]
    if isinstance(raw, dict) and isinstance(raw.get("items"), list):
        return [x for x in raw["items"] if isinstance(x, dict)]
    if isinstance(raw, dict):
        return [raw]
    return []


def _item_text(item: dict[str, Any]) -> str:
    parts = [
        item.get("title", ""),
        item.get("summary", ""),
        item.get("content", ""),
        item.get("transcript", ""),
        item.get("notes", ""),
    ]
    tags = _tags(item.get("tags"))
    if tags:
        parts.append("标签：" + "、".join(tags))
    if item.get("url"):
        parts.append(f"链接：{item['url']}")
    return "\n\n".join(str(p).strip() for p in parts if str(p).strip())


def _load_json_source(path: Path) -> list[dict[str, Any]]:
    chunks: list[dict[str, Any]] = []
    for i, item in enumerate(_json_items(path), start=1):
        text = _item_text(item)
        if not text:
            continue
        source_type = str(item.get("type") or "json_source")
        title = str(item.get("title") or item.get("id") or path.stem)
        source_id = str(item.get("id") or f"{path.stem}_{i}")
        kp = str(item.get("kp") or item.get("knowledge_point") or "")
        url = str(item.get("url") or "")
        for page, piece in enumerate(_split(text), start=1):
            chunks.append(_chunk_record(
                text=piece,
                source=title,
                chapter=source_type,
                kp=kp,
                page=page,
                source_type=source_type,
                url=url,
                tags=_tags(item.get("tags")),
                source_id=source_id,
            ))
    return chunks


def _load_text_source(path: Path) -> list[dict[str, Any]]:
    meta: dict[str, str] = {}
    body = _read_text(path)
    if path.suffix.lower() == ".md":
        meta, body = _front_matter(body)
    source = meta.get("source") or meta.get("title") or path.stem
    source_type = meta.get("type") or ("markdown" if path.suffix.lower() == ".md" else "text")
    chapter = meta.get("chapter") or ("extra markdown" if path.suffix.lower() == ".md" else "extra text")
    return [
        _chunk_record(
            text=piece,
            source=source,
            chapter=chapter,
            kp=meta.get("kp", ""),
            page=page,
            source_type=source_type,
            url=meta.get("url", ""),
            tags=_tags(meta.get("tags")),
            source_id=path.stem,
        )
        for page, piece in enumerate(_split(body), start=1)
    ]


def _load_extra_sources() -> list[dict[str, Any]]:
    if not EXTRA_SOURCES_DIR.exists():
        return []
    chunks: list[dict[str, Any]] = []
    for path in sorted(p for p in EXTRA_SOURCES_DIR.rglob("*") if p.is_file()):
        suffix = path.suffix.lower()
        if suffix not in SUPPORTED_EXTRA:
            continue
        try:
            if suffix == ".json":
                chunks.extend(_load_json_source(path))
            else:
                chunks.extend(_load_text_source(path))
        except Exception as exc:  # noqa: BLE001
            log.warning("extra source skipped: %s (%s)", path, exc)
    return chunks


def upload_meta_path(path: Path) -> Path:
    return path.with_name(f"{path.name}.meta.json")


def supported_upload_suffixes() -> set[str]:
    return set(SUPPORTED_UPLOAD)


def _upload_meta(path: Path) -> dict[str, Any]:
    meta_path = upload_meta_path(path)
    if not meta_path.exists():
        return {}
    try:
        data = json.loads(_read_text(meta_path))
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def source_status(meta: dict[str, Any], chunk_count: int = 0) -> str:
    """Normalize legacy sidecars into the public source lifecycle."""

    raw = str(meta.get("status") or "").strip().lower()
    if raw in {"", "active"}:
        return "indexed" if chunk_count > 0 else "processing"
    if raw == "empty":
        return "failed"
    return raw if raw in SOURCE_STATUSES else "inactive"


def source_is_active(meta: dict[str, Any]) -> bool:
    status = source_status(meta, int(meta.get("chunk_count") or 0))
    return meta.get("active", True) is not False and status not in INACTIVE_SOURCE_STATUSES


def source_is_selectable(meta: dict[str, Any], chunk_count: int) -> bool:
    return (
        source_is_active(meta)
        and chunk_count > 0
        and source_status(meta, chunk_count) in SELECTABLE_SOURCE_STATUSES
    )


def source_sha256(path: Path, meta: dict[str, Any] | None = None) -> str:
    stored = str((meta or {}).get("sha256") or "").strip().lower()
    return stored or hashlib.sha256(path.read_bytes()).hexdigest()


def _read_pdf_pages(path: Path) -> list[tuple[int, str]]:
    try:
        from pypdf import PdfReader
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError("PDF parsing requires pypdf. Run: pip install pypdf") from exc

    reader = PdfReader(str(path))
    pages: list[tuple[int, str]] = []
    for idx, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append((idx, text))
    return pages


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError("Word parsing requires python-docx. Run: pip install python-docx") from exc

    document = Document(str(path))
    blocks: list[str] = []
    blocks.extend(paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip())
    for table in document.tables:
        for row in table.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n\n".join(blocks)


def _load_uploaded_source(path: Path) -> list[dict[str, Any]]:
    meta = _upload_meta(path)
    user_id = str(meta.get("user_id") or "")
    source_id = str(meta.get("id") or path.stem)
    source = str(meta.get("title") or meta.get("filename") or path.stem)
    kp = str(meta.get("kp") or "")
    source_type = str(meta.get("source_type") or f"uploaded_{path.suffix.lower().lstrip('.')}")
    chapter = str(meta.get("filename") or "uploaded file")
    url = str(meta.get("url") or "")
    tags = _tags(meta.get("tags"))

    chunks: list[dict[str, Any]] = []
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        page_items = _read_pdf_pages(path)
        for page, text in page_items:
            for piece in _split(text):
                chunks.append(_chunk_record(
                    text=piece,
                    source=source,
                    chapter=chapter,
                    kp=kp,
                    page=page,
                    source_type=source_type,
                    url=url,
                    tags=tags,
                    source_id=source_id,
                    user_id=user_id,
                ))
        return chunks

    if suffix == ".docx":
        body = _read_docx(path)
    elif suffix in IMAGE_UPLOAD_SUFFIXES:
        ocr = meta.get("ocr") if isinstance(meta.get("ocr"), dict) else {}
        body = str(ocr.get("text") or "").strip()
        if not body:
            return []
        for page, piece in enumerate(_split(body), start=1):
            chunks.append(_chunk_record(
                text=piece,
                source=source,
                chapter=chapter,
                kp=kp,
                page=page,
                source_type=source_type,
                url=url,
                tags=tags,
                source_id=source_id,
                user_id=user_id,
            ))
        return chunks

    else:
        body = _read_text(path)
    if suffix == ".md":
        front, body = _front_matter(body)
        source = str(meta.get("title") or front.get("source") or front.get("title") or source)
        kp = str(meta.get("kp") or front.get("kp") or kp)
        source_type = str(meta.get("source_type") or front.get("type") or source_type)
        url = str(meta.get("url") or front.get("url") or url)
        tags = _tags(meta.get("tags") or front.get("tags"))

    for page, piece in enumerate(_split(body), start=1):
        chunks.append(_chunk_record(
            text=piece,
            source=source,
            chapter=chapter,
            kp=kp,
            page=page,
            source_type=source_type,
            url=url,
            tags=tags,
            source_id=source_id,
            user_id=user_id,
        ))
    return chunks


def _load_uploaded_sources() -> list[dict[str, Any]]:
    if not UPLOAD_SOURCES_DIR.exists():
        return []
    chunks: list[dict[str, Any]] = []
    seen_hashes: set[tuple[str, str]] = set()
    for path in sorted(p for p in UPLOAD_SOURCES_DIR.rglob("*") if p.is_file()):
        if path.name.endswith(".meta.json"):
            continue
        if path.suffix.lower() not in SUPPORTED_UPLOAD:
            continue
        try:
            meta = _upload_meta(path)
            if not source_is_active(meta):
                log.info("uploaded source quarantined, skipped: %s", path.name)
                continue
            digest = source_sha256(path, meta)
            owner_digest = (str(meta.get("user_id") or ""), digest)
            if owner_digest in seen_hashes:
                log.warning("duplicate uploaded source skipped: %s", path.name)
                continue
            seen_hashes.add(owner_digest)
            chunks.extend(_load_uploaded_source(path))
        except Exception as exc:  # noqa: BLE001
            log.warning("uploaded source skipped: %s (%s)", path, exc)
    return chunks


def chunks_for_source(source_id: str, user_id: str | None = None) -> list[dict[str, Any]]:
    return [
        chunk for chunk in load_all_chunks()
        if chunk.get("source_id") == source_id
        and (user_id is None or str(chunk.get("user_id") or "") == str(user_id))
    ]


def validate_source_scope(user_id: str, source_ids: list[str]) -> list[str]:
    """Return normalized source ids after ownership, activity, and index checks."""
    normalized = list(dict.fromkeys(
        str(source_id).strip() for source_id in source_ids if str(source_id).strip()
    ))
    if not normalized:
        return []

    metas: dict[str, dict[str, Any]] = {}
    if UPLOAD_SOURCES_DIR.exists():
        for meta_path in UPLOAD_SOURCES_DIR.rglob("*.meta.json"):
            try:
                meta = json.loads(_read_text(meta_path))
            except Exception:
                continue
            if isinstance(meta, dict) and meta.get("id"):
                metas[str(meta["id"])] = meta

    indexed = {str(chunk.get("source_id") or "") for chunk in load_all_chunks()}
    invalid = [
        source_id
        for source_id in normalized
        if source_id != COURSE_SCOPE_ID
        and (source_id not in metas
        or str(metas[source_id].get("user_id") or "") != user_id
        or not source_is_selectable(metas[source_id], 1 if source_id in indexed else 0))
    ]
    if invalid:
        raise ValueError("one or more selected sources are unavailable")
    return normalized


def load_all_chunks() -> list[dict[str, Any]]:
    return _load_seed_corpus() + _load_extra_sources() + _load_uploaded_sources()


def ingest_corpus(force: bool = False) -> int:
    """Build or refresh the vector store. Returns total chunk count."""
    store = get_store()
    manifest = _current_manifest()
    if store.count() > 0 and not force and _manifest_matches(manifest):
        log.info("knowledge base exists (%d chunks), manifest unchanged", store.count())
        return store.count()

    if store.count() > 0:
        log.info("knowledge base manifest changed, rebuilding vector store")
        clear = getattr(store, "clear", None)
        if callable(clear):
            clear()

    all_chunks = load_all_chunks()
    if not all_chunks:
        log.warning("no corpus chunks found, knowledge base not built")
        return 0

    vecs = embed_texts([c["text"] for c in all_chunks])
    store.upsert(all_chunks, vecs)
    _write_manifest(manifest)
    by_type: dict[str, int] = {}
    for chunk in all_chunks:
        key = str(chunk.get("source_type", "unknown"))
        by_type[key] = by_type.get(key, 0) + 1
    log.info("knowledge base built: %d chunks (%s)", len(all_chunks), by_type)
    return store.count()
